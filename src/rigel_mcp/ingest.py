"""Parsers that turn lesson files into citable chunks, plus library sync.

Library layout:  <library>/<course>/<lesson>/<any files, nested folders allowed>
Files placed directly in <library>/<course>/ belong to the lesson "_chung" (general).
"""

from __future__ import annotations

import hashlib
import json
import os
import re
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path

from .store import Store

CHUNK_LIMIT = 1500
CODE_WINDOW = 60
GENERAL_LESSON = "_chung"

TEXT_SUFFIXES = {".md", ".markdown", ".txt", ".rst"}
CODE_SUFFIXES = {".java", ".py", ".c", ".h", ".cpp", ".hpp", ".cs", ".js", ".ts", ".rs", ".go", ".sql", ".sh"}
KIND_BY_SUFFIX = {
    ".pdf": "pdf",
    ".pptx": "pptx",
    ".docx": "docx",
    ".ipynb": "notebook",
    **{s: "text" for s in TEXT_SUFFIXES},
    **{s: "code" for s in CODE_SUFFIXES},
}

SIGNATURE_RE = re.compile(
    r"^\s*(?:public|private|protected|static|abstract|final|class|interface|enum|def|function|fn|struct|impl|async)\b.*",
    re.MULTILINE,
)


class NoTextError(ValueError):
    """The file parsed fine but yielded no text (for example a scanned PDF)."""


@dataclass
class Chunk:
    locator: str
    heading: str
    text: str


@dataclass
class Parsed:
    chunks: list[Chunk] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


# ---- splitting -----------------------------------------------------------------------------


def _clean(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    text = re.sub(r"[ \t]+\n", "\n", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def _pieces(text: str, limit: int):
    """Yield (separator_before, piece), each piece at most `limit` characters."""
    first = True
    for para in re.split(r"\n\s*\n", text):
        sep = "" if first else "\n\n"
        first = False
        if len(para) <= limit:
            yield sep, para
            continue
        for i, line in enumerate(para.split("\n")):
            line_sep = sep if i == 0 else "\n"
            if len(line) <= limit:
                yield line_sep, line
            else:
                for j in range(0, len(line), limit):
                    yield (line_sep if j == 0 else "\n"), line[j : j + limit]


def split_text(text: str, limit: int = CHUNK_LIMIT) -> list[str]:
    """Split on paragraph, then line boundaries so that no part exceeds `limit` characters."""
    out: list[str] = []
    cur = ""
    for sep, piece in _pieces(_clean(text), limit):
        if cur and len(cur) + len(sep) + len(piece) > limit:
            out.append(cur)
            cur = piece
        else:
            cur = f"{cur}{sep}{piece}" if cur else piece
    if cur:
        out.append(cur)
    return out


def _first_line(text: str, width: int = 90) -> str:
    for line in text.splitlines():
        if line.strip():
            return line.strip()[:width]
    return ""


def _add(out: Parsed, base: str, heading: str, text: str) -> None:
    for n, part in enumerate(split_text(text), start=1):
        out.chunks.append(Chunk(base + (f" (phần {n})" if n > 1 else ""), heading, part))


# ---- parsers -------------------------------------------------------------------------------


def _toc_heading(toc: list, page: int) -> str:
    title = ""
    for _level, name, target in toc:
        if 0 < target <= page:
            title = name.strip()
    return title


def parse_pdf(path: Path) -> Parsed:
    import pymupdf

    out = Parsed()
    empty = 0
    with pymupdf.open(str(path)) as doc:
        toc = doc.get_toc(simple=True)
        for number, page in enumerate(doc, start=1):
            text = _clean(page.get_text("text"))
            if not text:
                empty += 1
                continue
            _add(out, f"trang {number}", _toc_heading(toc, number) or _first_line(text), text)
    if empty:
        out.warnings.append(f"{path.name}: {empty} trang không có text (có thể là bản scan, cần OCR)")
    return out


def parse_pptx(path: Path) -> Parsed:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    def shape_lines(shape, skip_id) -> list[str]:
        if getattr(shape, "shape_id", None) == skip_id:
            return []
        try:
            is_group = shape.shape_type == MSO_SHAPE_TYPE.GROUP
        except NotImplementedError:
            is_group = False
        if is_group:
            return [line for child in shape.shapes for line in shape_lines(child, skip_id)]
        lines: list[str] = []
        if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                t = p.text.strip()
                if t:
                    lines.append("  " * p.level + t)
        if getattr(shape, "has_table", False) and shape.has_table:
            for row in shape.table.rows:
                cells = [c.text.strip() for c in row.cells]
                if any(cells):
                    lines.append(" | ".join(cells))
        return lines

    out = Parsed()
    picture_only = 0
    for number, slide in enumerate(Presentation(str(path)).slides, start=1):
        title_shape = slide.shapes.title
        title = title_shape.text_frame.text.strip() if title_shape is not None and title_shape.has_text_frame else ""
        skip_id = title_shape.shape_id if title_shape is not None else None
        body = [line for shape in slide.shapes for line in shape_lines(shape, skip_id)]
        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        parts = ([title] if title else []) + body + ([f"Ghi chú: {notes}"] if notes else [])
        text = "\n".join(parts).strip()
        if not text:
            picture_only += 1
            continue
        heading = title or (body[0].strip()[:90] if body else "")
        _add(out, f"slide {number}", heading, text)
    if picture_only:
        out.warnings.append(f"{path.name}: {picture_only} slide chỉ có hình, không có text")
    return out


def parse_docx(path: Path) -> Parsed:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = docx.Document(str(path))
    out = Parsed()
    stack: list[str] = []
    lines: list[str] = []
    heading = ""

    def flush() -> None:
        text = "\n".join(lines).strip()
        lines.clear()
        if text:
            _add(out, f"§ {heading}" if heading else "đầu tài liệu", heading, text)

    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = Paragraph(child, document)
            text = para.text.strip()
            if not text:
                continue
            try:
                style = para.style.name if para.style is not None else ""
            except Exception:  # noqa: BLE001 - a broken style table must not stop ingestion
                style = ""
            m = re.match(r"(?i)^heading\s*(\d+)$", style or "")
            if m or (style or "").lower() == "title":
                flush()
                level = int(m.group(1)) if m else 1
                stack = stack[: level - 1] + [text]
                heading = " > ".join(stack)
            else:
                lines.append(text)
        elif tag == "tbl":
            for row in Table(child, document).rows:
                cells: list[str] = []
                for cell in row.cells:
                    value = cell.text.strip()
                    if not cells or cells[-1] != value:  # merged cells repeat their text
                        cells.append(value)
                if any(cells):
                    lines.append(" | ".join(cells))
    flush()
    return out


def parse_text(path: Path) -> Parsed:
    out = Parsed()
    stack: list[str] = []
    lines: list[str] = []
    heading = ""
    in_fence = False

    def flush() -> None:
        text = "\n".join(lines).strip()
        lines.clear()
        if text:
            _add(out, f"§ {heading}" if heading else "đầu file", heading, text)

    for line in path.read_text(encoding="utf-8", errors="replace").splitlines():
        if line.lstrip().startswith(("```", "~~~")):
            in_fence = not in_fence
        m = None if in_fence else re.match(r"^(#{1,6})\s+(.*\S)\s*$", line)
        if m:
            flush()
            stack = stack[: len(m.group(1)) - 1] + [m.group(2).strip()]
            heading = " > ".join(stack)
        else:
            lines.append(line)
    flush()
    return out


def parse_code(path: Path) -> Parsed:
    lines = path.read_text(encoding="utf-8", errors="replace").splitlines()
    out = Parsed()
    for start in range(0, len(lines), CODE_WINDOW):
        part = lines[start : start + CODE_WINDOW]
        text = "\n".join(part).rstrip()
        if not text.strip():
            continue
        m = SIGNATURE_RE.search(text)
        heading = m.group(0).strip()[:80] if m else path.name
        out.chunks.append(Chunk(f"dòng {start + 1}–{start + len(part)}", heading, text))
    return out


def parse_ipynb(path: Path) -> Parsed:
    notebook = json.loads(path.read_text(encoding="utf-8"))
    out = Parsed()
    for number, cell in enumerate(notebook.get("cells", []), start=1):
        source = cell.get("source", "")
        source = ("".join(source) if isinstance(source, list) else str(source)).strip()
        if not source:
            continue
        kind = cell.get("cell_type")
        if kind == "markdown":
            m = re.search(r"^#{1,6}\s+(.+)$", source, re.MULTILINE)
            heading, text = (m.group(1).strip() if m else _first_line(source)), source
        elif kind == "code":
            heading, text = "code: " + _first_line(source, 60), f"```\n{source}\n```"
        else:
            continue
        _add(out, f"cell {number}", heading, text)
    return out


PARSERS = {
    "pdf": parse_pdf,
    "pptx": parse_pptx,
    "docx": parse_docx,
    "text": parse_text,
    "code": parse_code,
    "notebook": parse_ipynb,
}


def parse_file(path: Path) -> tuple[str, Parsed]:
    kind = KIND_BY_SUFFIX[path.suffix.lower()]
    return kind, PARSERS[kind](path)


# ---- ingest and sync -----------------------------------------------------------------------


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1 << 20), b""):
            h.update(block)
    return h.hexdigest()


@dataclass
class IngestResult:
    doc_id: int
    created: bool
    n_chunks: int
    warnings: list[str]


def abs_posix(path: Path) -> str:
    """Absolute path with forward slashes, as stored in docs.abs_path and cited in Vega cards."""
    return Path(os.path.abspath(path)).as_posix()


def ingest_file(
    store: Store,
    path: Path,
    course: str,
    lesson: str,
    *,
    origin: str = "library",
    rel_path: str | None = None,
    sha256: str | None = None,
    size: int | None = None,
    mtime: float | None = None,
    abs_path: str | None = None,
) -> IngestResult:
    kind, parsed = parse_file(path)
    if not parsed.chunks:
        raise NoTextError(f"{path.name}: " + ("; ".join(parsed.warnings) or "không trích được text"))
    if size is None or mtime is None:
        st = path.stat()
        size, mtime = st.st_size, st.st_mtime
    doc_id, created = store.upsert_doc(
        course=course,
        lesson=lesson,
        rel_path=rel_path or path.name,
        kind=kind,
        sha256=sha256 or sha256_file(path),
        origin=origin,
        chunks=parsed.chunks,
        size=size,
        mtime=mtime,
        abs_path=abs_path or abs_posix(path),
    )
    return IngestResult(doc_id, created, len(parsed.chunks), parsed.warnings)


@dataclass
class SyncReport:
    added: list[str] = field(default_factory=list)
    updated: list[str] = field(default_factory=list)
    removed: list[str] = field(default_factory=list)
    unchanged: int = 0
    chunks: int = 0
    warnings: list[str] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)
    mode: str = "folder"  # "folder" (classic library) or "map" (library.json)
    source: str = ""  # library folder or map file that was synced
    kept: int = 0  # docs kept because their source folder could not be read
    discovery: dict[str, int] = field(default_factory=dict)  # how files were found: method -> count
    notes: list[str] = field(default_factory=list)  # reminders for the learner

    def as_dict(self) -> dict:
        return {
            "added": self.added,
            "updated": self.updated,
            "removed": self.removed,
            "unchanged": self.unchanged,
            "chunks_written": self.chunks,
            "warnings": self.warnings,
            "errors": self.errors,
            "mode": self.mode,
            "source": self.source,
            "kept": self.kept,
            "discovery": self.discovery,
            "notes": self.notes,
        }

    def render(self) -> str:
        lines = [
            f"Đồng bộ xong: +{len(self.added)} mới, ~{len(self.updated)} cập nhật, "
            f"-{len(self.removed)} gỡ, {self.unchanged} không đổi (ghi {self.chunks} đoạn)."
        ]
        for label, items in (("Mới", self.added), ("Cập nhật", self.updated), ("Đã gỡ", self.removed)):
            if items:
                shown = ", ".join(items[:10]) + (f" … (+{len(items) - 10})" if len(items) > 10 else "")
                lines.append(f"{label}: {shown}")
        if self.mode == "map":
            lines.append(f"Bản đồ: {self.source}")
            labels = (
                ("smartdrive", "chỉ mục SmartDrive"),
                ("listing", "liệt kê thư mục"),
                ("walk", "duyệt thư mục có giới hạn"),
                ("file", "file đơn"),
            )
            found = ", ".join(f"{text} {self.discovery[key]}" for key, text in labels if self.discovery.get(key))
            lines.append("Cách tìm file: " + (found or "không có file nào") + ".")
        if self.kept:
            lines.append(f"Giữ nguyên {self.kept} tài liệu đã nạp vì không truy cập được thư mục nguồn.")
        lines += self.notes
        lines += [f"Cảnh báo: {w}" for w in self.warnings[:10]]
        if len(self.warnings) > 10:
            lines.append(f"… còn {len(self.warnings) - 10} cảnh báo nữa (xem structured_content.warnings).")
        lines += [f"Lỗi: {e}" for e in self.errors[:10]]
        return "\n".join(lines)


def _hidden(name: str) -> bool:
    return name.startswith((".", "~$"))


def iter_library(root: Path):
    """Yield (course, lesson, rel_path, file) for every file under <root>/<course>/..."""
    for course_dir in sorted(root.iterdir()):
        if not course_dir.is_dir() or _hidden(course_dir.name):
            continue
        for entry in sorted(course_dir.iterdir()):
            if _hidden(entry.name):
                continue
            if entry.is_dir():
                for f in sorted(entry.rglob("*")):
                    rel = f.relative_to(entry)
                    if f.is_file() and not any(_hidden(part) for part in rel.parts):
                        yield course_dir.name, entry.name, rel.as_posix(), f
            elif entry.is_file():
                yield course_dir.name, GENERAL_LESSON, entry.name, entry


def sync_one(
    store: Store,
    rep: SyncReport,
    known: dict[tuple[str, str, str], dict],
    *,
    f: Path,
    course: str,
    lesson: str,
    rel: str,
    origin: str,
) -> None:
    """Bring one file up to date, reading it as little as possible.

    Same size and mtime as recorded: skip without reading. Otherwise hash it once; an unchanged hash only
    refreshes size/mtime (no re-parse), a changed one is re-ingested under the same doc id.
    """
    label = f"{course}/{lesson}/{rel}"
    try:
        st = f.stat()
    except OSError as e:
        rep.errors.append(f"{label}: {type(e).__name__}: {e}")
        return
    size, mtime, abs_path = st.st_size, st.st_mtime, abs_posix(f)
    try:
        doc = known.get((course, lesson, rel))
        if doc and doc["size"] == size and doc["mtime"] == mtime:
            if doc["origin"] != origin or doc["abs_path"] != abs_path:
                store.touch_doc(doc["id"], size=size, mtime=mtime, abs_path=abs_path, origin=origin)
            rep.unchanged += 1
            return
        sha = sha256_file(f)
        if doc and doc["sha256"] == sha:
            store.touch_doc(doc["id"], size=size, mtime=mtime, abs_path=abs_path, origin=origin)
            rep.unchanged += 1
            return
        result = ingest_file(
            store, f, course, lesson, origin=origin, rel_path=rel, sha256=sha, size=size, mtime=mtime, abs_path=abs_path
        )
    except NoTextError as e:
        rep.warnings.append(str(e))
        return
    except Exception as e:  # noqa: BLE001 - one bad file must not stop the sync
        rep.errors.append(f"{label}: {type(e).__name__}: {e}")
        return
    (rep.added if result.created else rep.updated).append(label)
    rep.chunks += result.n_chunks
    rep.warnings += result.warnings


def sync_library(store: Store, root: Path) -> SyncReport:
    """Classic mode: <root>/<course>/<lesson>/<files>. Only docs with origin "library" are ever removed."""
    rep = SyncReport(mode="folder", source=str(root))
    seen: set[tuple[str, str, str]] = set()
    unsupported: Counter[str] = Counter()
    known = store.docs_index()

    for f in sorted(root.iterdir()):
        if f.is_file() and not _hidden(f.name) and f.suffix.lower() in KIND_BY_SUFFIX:
            rep.warnings.append(f"{f.name}: nằm ngay ở gốc thư viện nên bị bỏ qua — đặt vào <môn>/<bài>/")

    for course, lesson, rel, f in iter_library(root):
        if f.suffix.lower() not in KIND_BY_SUFFIX:
            unsupported[f.suffix.lower() or "(không đuôi)"] += 1
            continue
        seen.add((course, lesson, rel))
        sync_one(store, rep, known, f=f, course=course, lesson=lesson, rel=rel, origin="library")

    for doc in store.list_docs("library"):
        if (doc["course"], doc["lesson"], doc["rel_path"]) not in seen:
            store.delete_doc(doc["id"])
            rep.removed.append(f"{doc['course']}/{doc['lesson']}/{doc['rel_path']}")

    if unsupported:
        kinds = ", ".join(f"{ext} ×{n}" for ext, n in unsupported.most_common(6))
        rep.warnings.append(f"bỏ qua file định dạng chưa hỗ trợ: {kinds}")
    return rep
