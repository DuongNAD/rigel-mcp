import json
from dataclasses import dataclass
from pathlib import Path

import docx as pydocx
import pymupdf
import pytest
from pptx import Presentation

from rigel_mcp.config import Config

LESSON = ("PRJ301", "L03-OOP")


@pytest.fixture(autouse=True)
def isolated_env(monkeypatch):
    """No test may read the learner's real map, hub or SmartDrive index, whatever the machine's environment says."""
    for name in (
        "RIGEL_LIBRARY_MAP",
        "RIGEL_HUB",
        "RIGEL_SMARTDRIVE_CMD",
        "RIGEL_SMARTDRIVE_ROOT",
    ):
        monkeypatch.delenv(name, raising=False)
    monkeypatch.setenv("RIGEL_SMARTDRIVE", "off")


@pytest.fixture
def library(tmp_path):
    root = tmp_path / "lib"
    lesson = root / LESSON[0] / LESSON[1]
    lesson.mkdir(parents=True)

    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[1])
    s1.shapes.title.text = "Polymorphism"
    s1.placeholders[1].text_frame.text = "Same interface, many forms\nOverriding vs overloading"
    s1.notes_slide.notes_text_frame.text = "Nhấn mạnh dynamic dispatch"
    s2 = prs.slides.add_slide(prs.slide_layouts[1])
    s2.shapes.title.text = "Encapsulation"
    s2.placeholders[1].text_frame.text = "Hide state behind methods\nprivate fields, public getters"
    prs.save(lesson / "Lecture3.pptx")

    doc = pydocx.Document()
    doc.add_heading("Lập trình hướng đối tượng", 1)
    doc.add_paragraph("Bốn tính chất: đóng gói, kế thừa, đa hình, trừu tượng.")
    doc.add_heading("Đa hình", 2)
    doc.add_paragraph("Đa hình cho phép một lời gọi phương thức cho hành vi khác nhau tùy kiểu thực tế của đối tượng.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Overload", "Cùng tên, khác tham số"
    table.cell(1, 0).text, table.cell(1, 1).text = "Override", "Lớp con ghi đè lớp cha"
    doc.save(lesson / "Handout.docx")

    pdf = pymupdf.open()
    for lines in (
        ["Inheritance lets a class reuse", "fields and methods of a parent."],
        ["Abstract classes cannot be instantiated;", "interfaces define contracts."],
    ):
        pdf.new_page().insert_text((72, 72), "\n".join(lines))
    pdf.set_toc([[1, "Inheritance", 1], [1, "Abstraction", 2]])
    pdf.save(str(lesson / "Slides.pdf"))
    pdf.close()

    (lesson / "notes.md").write_text(
        "# Đa hình\n\nOverriding ghi đè phương thức của lớp cha.\n\n"
        "```python\n# comment này không phải heading\nprint('x')\n```\n\n"
        "## Overload\n\nCùng tên, khác tham số.\n",
        encoding="utf-8",
    )
    (lesson / "Shape.java").write_text(
        "public class Shape {\n    public double area() { return 0; }\n}\n", encoding="utf-8"
    )
    return root


@pytest.fixture
def cfg(tmp_path, library):
    return Config(home=tmp_path / "home", library=library, notes=tmp_path / "notes")


# ---- library map fixtures (0.2.0) ----------------------------------------------------------

HUB_FILES = {
    # course GCI: root "subfolders" (03_Materials), "file" (ROADMAP.md) and "single" (06_Notes)
    "GCI_Subject/03_Materials/overview.md": "# Tổng quan\n\nGiới thiệu khóa học GCI về phân tích dữ liệu.\n",
    "GCI_Subject/03_Materials/01_Intro/slides.md": "# Giới thiệu\n\nPython là ngôn ngữ lập trình đa dụng.\n",
    "GCI_Subject/03_Materials/01_Intro/code/hello.py": "def hello():\n    return 'xin chào'\n",
    "GCI_Subject/03_Materials/01_Intro/__pycache__/junk.py": "x = 1\n",
    "GCI_Subject/03_Materials/01_Intro/.hidden/secret.md": "# bí mật\n\nkhông được nạp\n",
    "GCI_Subject/03_Materials/01_Intro/~$lock.md": "# khóa\n\nfile khóa của Office\n",
    "GCI_Subject/03_Materials/02_NumPy/notes.md": (
        "# NumPy\n\nMảng ndarray và broadcasting cho phép tính toán hàng loạt trên dữ liệu số "
        "mà không cần vòng lặp Python.\n"
    ),
    "GCI_Subject/03_Materials/02_NumPy/dap_an/answers.md": "# Đáp án\n\nLời giải bài tập\n",
    "GCI_Subject/03_Materials/node_modules/pkg.md": "# pkg\n\nthư viện bên thứ ba\n",
    "GCI_Subject/roadmap/ROADMAP.md": "# Lộ trình\n\nTuần 1 học Python, tuần 2 học NumPy.\n",
    "GCI_Subject/06_Notes/lecture.txt": "Buổi giảng về pandas và dataframe.\n",
    "GCI_Subject/06_Notes/extra/more.md": "# Thêm\n\nGhi chú bổ sung về dataframe.\n",
    # course ML: only subfolders
    "Machine_Learning/02_Supervised_Learning/01_Regression/regression.py": (
        "def fit(xs, ys):\n    return sum(xs) / len(xs)\n"
    ),
    "Machine_Learning/02_Supervised_Learning/02_Classification/notes.md": "# Phân loại\n\nLogistic regression phân loại nhị phân.\n",
}

# What a correct sync of HUB_FILES produces: (course, lesson, rel_path)
HUB_DOCS = {
    ("GCI", "_chung", "overview.md"),
    ("GCI", "01_Intro", "slides.md"),
    ("GCI", "01_Intro", "code/hello.py"),
    ("GCI", "02_NumPy", "notes.md"),
    ("GCI", "_lo_trinh", "ROADMAP.md"),
    ("GCI", "Bai_giang", "lecture.txt"),
    ("GCI", "Bai_giang", "extra/more.md"),
    ("ML", "01_Regression", "regression.py"),
    ("ML", "02_Classification", "notes.md"),
}


def write_tree(root: Path, files: dict[str, str]) -> None:
    for rel, text in files.items():
        f = root / rel
        f.parent.mkdir(parents=True, exist_ok=True)
        f.write_text(text, encoding="utf-8")


@dataclass
class Hub:
    root: Path  # the "hub" (stands in for D:/02_Learning_Knowledge)
    tmp: Path

    @property
    def materials(self) -> Path:
        return self.root / "GCI_Subject" / "03_Materials"

    @property
    def map_file(self) -> Path:
        return self.tmp / "home" / "library.json"

    @property
    def expected(self) -> set[tuple[str, str, str]]:
        return set(HUB_DOCS)

    def map_dict(self) -> dict:
        r = self.root.as_posix()
        return {
            "v": 1,
            "hub": r,
            "exclude": ["__pycache__", "node_modules"],
            "maxFileMB": 40,
            "courses": [
                {
                    "name": "GCI",
                    "subject": "GCI_Subject",
                    "aliases": ["GCI World", "gci"],
                    "roots": [
                        {"path": f"{r}/GCI_Subject/03_Materials", "lessons": "subfolders"},
                        {"path": f"{r}/GCI_Subject/roadmap/ROADMAP.md", "lessons": "file", "lesson": "_lo_trinh"},
                        {"path": f"{r}/GCI_Subject/06_Notes", "lessons": "single", "lesson": "Bai_giang"},
                    ],
                    "exclude": ["dap_an"],
                },
                {
                    "name": "ML",
                    "subject": "Machine_Learning",
                    "aliases": ["Học máy", "ml"],
                    "roots": [{"path": f"{r}/Machine_Learning/02_Supervised_Learning", "lessons": "subfolders"}],
                },
            ],
        }

    def write_map(self, data: dict | None = None) -> Path:
        self.map_file.parent.mkdir(parents=True, exist_ok=True)
        self.map_file.write_text(json.dumps(data or self.map_dict(), ensure_ascii=False, indent=2), encoding="utf-8")
        return self.map_file


@pytest.fixture
def hub(tmp_path):
    root = tmp_path / "hub"
    write_tree(root, HUB_FILES)
    return Hub(root=root, tmp=tmp_path)


@pytest.fixture
def map_cfg(tmp_path, hub):
    """Server config in map mode: the classic library folder does not even exist."""
    hub.write_map()
    return Config(
        home=tmp_path / "home",
        library=tmp_path / "no-library",
        notes=tmp_path / "notes",
        map_path=hub.map_file,
        hub=hub.root,
    )
